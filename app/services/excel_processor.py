from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from typing import Any
from datetime import date, datetime

from openpyxl import load_workbook

STUDENT_ID_RE = re.compile(r"^[A-Z]{2}\d{7}$")

class WorkbookValidationError(ValueError):
    pass

def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()

def _student_id(value: Any, row_num: int) -> str:
    result = _cell_text(value).upper()
    if not STUDENT_ID_RE.fullmatch(result):
        raise WorkbookValidationError(
            f"Row {row_num}: expected student ID with two English letters and seven digits (e.g. AB1234567)."
        )
    return result

def _decimal(value: Any, row_num: int, column_name: str) -> Decimal:
    if value is None or str(value).strip() == "":
        raise WorkbookValidationError(f"Row {row_num}: {column_name} is required.")
    try:
        if isinstance(value, str):
            cleaned = value.strip().replace(" ", "").replace(",", "")
            if cleaned.endswith("%"):
                cleaned = cleaned[:-1]
            result = Decimal(cleaned)
        else:
            result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise WorkbookValidationError(f"Row {row_num}: {column_name} must be numeric.")
    if not result.is_finite():
        raise WorkbookValidationError(f"Row {row_num}: {column_name} must be a finite number.")
    return result

def parse_credentials(file_bytes: bytes) -> dict[str, dict[str, str]]:
    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookValidationError("Could not read this Excel workbook. Please upload a valid .xlsx file.") from exc
    if not wb.worksheets:
        raise WorkbookValidationError("The workbook has no worksheets.")
    ws = wb.worksheets[0]
    result: dict[str, dict[str, str]] = {}
    try:
        for row_num, row in enumerate(ws.iter_rows(min_row=2, min_col=1, max_col=2, values_only=True), start=2):
            login_value, password_value = row
            login = _cell_text(login_value)
            password = _cell_text(password_value).upper()
            if not login and not password:
                continue
            if not login:
                raise WorkbookValidationError(f"Row {row_num}: HEMIS login (column A) is empty.")
            student_id = _student_id(password_value, row_num)
            if student_id in result:
                raise WorkbookValidationError(f"Row {row_num}: duplicate student ID {student_id}.")
            result[student_id] = {"login": login, "password": password}
    finally:
        wb.close()
    if not result:
        raise WorkbookValidationError("No credential records found from row 2 onward.")
    return result

def parse_contracts(file_bytes: bytes) -> dict[str, dict[str, int | float]]:
    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookValidationError("Could not read this Excel workbook. Please upload a valid .xlsx file.") from exc
    if not wb.worksheets:
        raise WorkbookValidationError("The workbook has no worksheets.")
    ws = wb.worksheets[0]
    raw_update_date = ws["A1"].value

    if isinstance(raw_update_date, datetime):
        update_date = raw_update_date.date().isoformat()
    elif isinstance(raw_update_date, date):
        update_date = raw_update_date.isoformat()
    elif raw_update_date is not None:
        update_date = str(raw_update_date).strip()
    else:
        raise WorkbookValidationError(
            "Cell A1 must contain the contract data update date."
        )
    result: dict[str, dict[str, int | float]] = {}
    try:
        for row_num, row in enumerate(ws.iter_rows(min_row=2, min_col=1, max_col=4, values_only=True), start=2):
            student_value, total_value, paid_value, percentage_value = row
            if all(value is None or str(value).strip() == "" for value in row):
                continue
            student_id = _student_id(student_value, row_num)
            if student_id in result:
                raise WorkbookValidationError(f"Row {row_num}: duplicate student ID {student_id}.")
            total = _decimal(total_value, row_num, "total contract amount")
            paid = _decimal(paid_value, row_num, "paid amount")
            supplied_percentage = _decimal(
                percentage_value, row_num, "payment percentage"
            )

            # Excel percentage-formatted numeric cells store 25.12% as 0.2512.
            if not isinstance(percentage_value, str) and supplied_percentage <= 1:
                supplied_percentage *= Decimal("100")
            if total <= 0:
                raise WorkbookValidationError(f"Row {row_num}: total contract amount must be greater than zero.")
            if paid <= 0:
                raise WorkbookValidationError(
                    f"Row {row_num}: paid amount must be greater than zero."
                )
            calculated_percentage = paid / total * Decimal("100")
            remaining = total - paid
            result[student_id] = {
                "total_amount": _number(total),
                "paid_amount": _number(paid),
                "remaining_amount": _number(remaining),
                "payment_percentage": float(supplied_percentage),
            }
    finally:
        wb.close()
    if not result:
        raise WorkbookValidationError("No contract records found from row 2 onward.")
    return {
        "update_date": update_date,
        "students": result,
    }

def _number(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)
    return float(value)
